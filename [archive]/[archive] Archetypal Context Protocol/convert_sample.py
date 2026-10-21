import docx
import os

# Convert the Hero archetype as a sample
base_path = r"C:\Users\User\Documents\0000000\Archetypal Context Protocol"
input_file = os.path.join(base_path, "[MSC] Modern Spiritual and Cultural", "[HJ] Hero's Journey", "[HJ] Hero.docx")
output_file = os.path.join(base_path, "[HJ] Hero_sample.txt")

try:
    # Load the document
    doc = docx.Document(input_file)

    # Extract all text
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)

    # Also extract text from tables if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)

    # Write to text file
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write('\n'.join(full_text))

    print(f"Successfully converted: {os.path.basename(input_file)}")
    print(f"Output saved to: {output_file}")

except Exception as e:
    print(f"Error converting file: {e}")
