import os
import docx
from pathlib import Path

base_path = r"C:\Users\User\Documents\0000000\Archetypal Context Protocol"

# Create output directory for converted files
output_dir = os.path.join(base_path, "converted_txt")
os.makedirs(output_dir, exist_ok=True)

def convert_docx_to_txt(docx_path, txt_path):
    """Convert a single .docx file to .txt"""
    try:
        # Load the document
        doc = docx.Document(docx_path)

        # Extract all text from paragraphs
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)

        # Also extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)

        # Write to text file
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(full_text))

        return True
    except Exception as e:
        print(f"  ERROR: {e}")
        return False

# Find and convert all .docx files
print("Searching for .docx files...")
converted_count = 0
failed_count = 0

for root, dirs, files in os.walk(base_path):
    for file in files:
        if file.endswith('.docx') and not file.startswith('~'):  # Skip temp files
            input_path = os.path.join(root, file)

            # Create relative path for organization
            rel_path = os.path.relpath(input_path, base_path)
            rel_dir = os.path.dirname(rel_path)

            # Create corresponding output directory
            output_subdir = os.path.join(output_dir, rel_dir)
            os.makedirs(output_subdir, exist_ok=True)

            # Create output filename
            output_filename = file.replace('.docx', '.txt')
            output_path = os.path.join(output_subdir, output_filename)

            print(f"\nConverting: {rel_path}")
            if convert_docx_to_txt(input_path, output_path):
                print(f"  -> {output_path}")
                converted_count += 1
            else:
                failed_count += 1

print(f"\n{'='*60}")
print(f"Conversion complete!")
print(f"Successfully converted: {converted_count} files")
print(f"Failed: {failed_count} files")
print(f"Output directory: {output_dir}")
